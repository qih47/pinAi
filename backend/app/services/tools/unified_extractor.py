"""
CAKRA AI — Unified Document Extractor
======================================
Engine ekstraksi dokumen multi-format terpusat.
Menggabungkan logika dari document_intelligence.py, peraturan_service.py,
dan vision_service.py menjadi satu antarmuka tunggal.

Hierarchy Cache:
  L1: RAM (_EXTRACTOR_CACHE) — sub-milidetik
  L2: Disk (CACHE_DIR/*.json) — antar restart

Formats Didukung:
  PDF (digital), PDF (scan/OCR), DOCX, XLSX, CSV, TXT, gambar (JPG/PNG)
"""

import os
import json
import base64
import logging
import asyncio
import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple
from dataclasses import dataclass, field, asdict
from concurrent.futures import ThreadPoolExecutor

logger = logging.getLogger("CAKRA_UNIFIED_EXTRACTOR")

# ─── Konstanta ──────────────────────────────────────────────────────────────
CACHE_DIR = Path("/home/qisthi/pinAi/backend/uploads/.ocr_cache")
CACHE_DIR.mkdir(parents=True, exist_ok=True)

MAX_WORKERS = 8          # Thread pool untuk OCR paralel
OCR_SCAN_THRESHOLD = 40  # Karakter per halaman, di bawah ini = scan
XLSX_MAX_ROWS = 500      # Batas baris tabel yang diekstrak ke teks


# ─── Data Model ─────────────────────────────────────────────────────────────
@dataclass
class ExtractedPage:
    """Representasi satu halaman dokumen yang sudah diekstrak."""
    page_num: int       # 0-indexed
    text: str           # Teks bersih
    is_scan: bool       # True jika halaman ini hasil OCR
    base64_image: str = ""  # Kosong jika tidak di-render


@dataclass
class ExtractedDocument:
    """Hasil ekstraksi lengkap satu dokumen."""
    file_path: str
    file_type: str                        # "pdf", "docx", "xlsx", "txt", "image"
    total_pages: int
    is_scanned: bool                      # True jika mayoritas halaman adalah scan
    full_text: str                        # Gabungan seluruh teks
    pages: List[ExtractedPage] = field(default_factory=list)
    base64_images: List[str] = field(default_factory=list)
    extracted_at: str = field(default_factory=lambda: datetime.datetime.now().isoformat())
    extraction_seconds: float = 0.0

    def to_text_map(self) -> List[Dict[str, Any]]:
        """Konversi ke format text_map yang dipakai document_intelligence."""
        return [{"page_num": p.page_num, "text": p.text} for p in self.pages]


# ─── L1 RAM Cache ───────────────────────────────────────────────────────────
_EXTRACTOR_CACHE: Dict[str, ExtractedDocument] = {}


def _get_l1(cache_key: str) -> Optional[ExtractedDocument]:
    return _EXTRACTOR_CACHE.get(cache_key)


def _set_l1(cache_key: str, doc: ExtractedDocument) -> None:
    _EXTRACTOR_CACHE[cache_key] = doc


# ─── L2 Disk Cache ──────────────────────────────────────────────────────────
def _get_l2(cache_key: str) -> Optional[ExtractedDocument]:
    safe = cache_key.replace("/", "_").replace("\\", "_")
    cache_file = CACHE_DIR / f"{safe}.extract.json"
    if not cache_file.exists():
        return None
    try:
        data = json.loads(cache_file.read_text(encoding="utf-8"))
        pages = [ExtractedPage(**p) for p in data.pop("pages", [])]
        doc = ExtractedDocument(**data, pages=pages)
        _set_l1(cache_key, doc)
        return doc
    except Exception as e:
        logger.warning(f"[EXTRACTOR] L2 cache baca gagal: {e}")
        return None


def _set_l2(cache_key: str, doc: ExtractedDocument) -> None:
    safe = cache_key.replace("/", "_").replace("\\", "_")
    cache_file = CACHE_DIR / f"{safe}.extract.json"
    try:
        data = asdict(doc)
        cache_file.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception as e:
        logger.warning(f"[EXTRACTOR] L2 cache tulis gagal: {e}")


# ─── Main Entry Point ────────────────────────────────────────────────────────
async def extract_document(
    file_path: str,
    cache_key: Optional[str] = None,
    render_images: bool = False,
    use_cache: bool = True,
) -> ExtractedDocument:
    """
    Ekstrak dokumen dari file_path secara async.

    Args:
        file_path: Path absolut ke file.
        cache_key: Key untuk cache (default: file_path).
        render_images: Jika True, render setiap halaman PDF sebagai base64 PNG.
        use_cache: Jika False, paksa ekstrak ulang meski ada cache.

    Returns:
        ExtractedDocument
    """
    key = cache_key or file_path

    if use_cache:
        # L1 check
        hit = _get_l1(key)
        if hit:
            logger.info(f"[EXTRACTOR] ⚡ L1 Cache HIT: {os.path.basename(file_path)}")
            return hit
        # L2 check
        hit = _get_l2(key)
        if hit:
            logger.info(f"[EXTRACTOR] ⚡ L2 Cache HIT: {os.path.basename(file_path)}")
            return hit

    if not os.path.exists(file_path):
        raise FileNotFoundError(f"[EXTRACTOR] File tidak ditemukan: {file_path}")

    ext = Path(file_path).suffix.lower()
    t0 = datetime.datetime.now()

    if ext == ".pdf":
        doc = await _extract_pdf(file_path, render_images=render_images)
    elif ext in (".docx", ".doc"):
        doc = await _extract_docx(file_path)
    elif ext in (".xlsx", ".xls"):
        doc = await _extract_xlsx(file_path)
    elif ext == ".csv":
        doc = await _extract_csv(file_path)
    elif ext in (".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp"):
        doc = await _extract_image(file_path)
    else:
        doc = await _extract_plain_text(file_path)

    doc.file_path = file_path
    doc.extraction_seconds = (datetime.datetime.now() - t0).total_seconds()
    logger.info(
        f"[EXTRACTOR] ✅ Selesai: {os.path.basename(file_path)} | "
        f"{doc.total_pages} halaman | {doc.extraction_seconds:.2f}s"
    )

    _set_l1(key, doc)
    _set_l2(key, doc)
    return doc


# ─── PDF Extractor ───────────────────────────────────────────────────────────
def _process_pdf_page_sync(file_path: str, page_num: int, render: bool) -> ExtractedPage:
    """
    Worker thread-safe: ekstrak satu halaman PDF.
    Digabungkan dari process_single_page_threadsafe (document_intelligence.py).
    """
    import fitz
    doc = fitz.open(file_path)
    page = doc.load_page(page_num)

    text = page.get_text("text").strip()
    is_scan = len(text) < OCR_SCAN_THRESHOLD
    b64_img = ""

    if render or is_scan:
        pix = page.get_pixmap(dpi=150)
        img_bytes = pix.tobytes("png")
        b64_img = base64.b64encode(img_bytes).decode("utf-8")

        if is_scan:
            try:
                import pytesseract
                from PIL import Image
                import io
                img = Image.open(io.BytesIO(img_bytes))
                ocr_text = pytesseract.image_to_string(img, lang="ind+eng")
                if ocr_text.strip():
                    text = ocr_text.strip()
            except Exception as e:
                logger.warning(f"[EXTRACTOR] OCR halaman {page_num+1} gagal: {e}")

    doc.close()
    return ExtractedPage(page_num=page_num, text=text, is_scan=is_scan, base64_image=b64_img)


async def _extract_pdf(file_path: str, render_images: bool = False) -> ExtractedDocument:
    """Ekstrak seluruh halaman PDF secara paralel."""
    import fitz
    doc_meta = fitz.open(file_path)
    total_pages = len(doc_meta)
    doc_meta.close()

    def run_parallel():
        with ThreadPoolExecutor(max_workers=MAX_WORKERS) as ex:
            futures = [
                ex.submit(_process_pdf_page_sync, file_path, p, render_images)
                for p in range(total_pages)
            ]
            return [f.result() for f in futures]

    pages: List[ExtractedPage] = await asyncio.to_thread(run_parallel)
    pages.sort(key=lambda p: p.page_num)

    full_text = "\n".join(p.text for p in pages if p.text)
    b64_images = [p.base64_image for p in pages if p.base64_image]
    is_scanned = sum(1 for p in pages if p.is_scan) > total_pages * 0.5

    return ExtractedDocument(
        file_path=file_path,
        file_type="pdf",
        total_pages=total_pages,
        is_scanned=is_scanned,
        full_text=full_text,
        pages=pages,
        base64_images=b64_images,
    )


# ─── DOCX Extractor ──────────────────────────────────────────────────────────
async def _extract_docx(file_path: str) -> ExtractedDocument:
    """Ekstrak teks dari file Word (.docx)."""
    def read_docx():
        try:
            from docx import Document
            doc = Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n".join(paragraphs)
        except ImportError:
            return "[python-docx tidak terinstal — ekstraksi DOCX gagal]"
        except Exception as e:
            return f"[Gagal baca DOCX: {e}]"

    text = await asyncio.to_thread(read_docx)
    page = ExtractedPage(page_num=0, text=text, is_scan=False)

    return ExtractedDocument(
        file_path=file_path,
        file_type="docx",
        total_pages=1,
        is_scanned=False,
        full_text=text,
        pages=[page],
    )


# ─── XLSX Extractor ──────────────────────────────────────────────────────────
async def _extract_xlsx(file_path: str) -> ExtractedDocument:
    """Ekstrak data tabel dari file Excel (.xlsx)."""
    def read_xlsx():
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            parts = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                parts.append(f"[Sheet: {sheet_name}]")
                row_count = 0
                for row in ws.iter_rows(values_only=True):
                    if row_count >= XLSX_MAX_ROWS:
                        parts.append(f"... (terpotong setelah {XLSX_MAX_ROWS} baris)")
                        break
                    cells = [str(c) if c is not None else "" for c in row]
                    if any(cells):
                        parts.append("\t".join(cells))
                    row_count += 1
            return "\n".join(parts)
        except ImportError:
            return "[openpyxl tidak terinstal — ekstraksi XLSX gagal]"
        except Exception as e:
            return f"[Gagal baca XLSX: {e}]"

    text = await asyncio.to_thread(read_xlsx)
    page = ExtractedPage(page_num=0, text=text, is_scan=False)

    return ExtractedDocument(
        file_path=file_path,
        file_type="xlsx",
        total_pages=1,
        is_scanned=False,
        full_text=text,
        pages=[page],
    )


# ─── CSV Extractor ───────────────────────────────────────────────────────────
async def _extract_csv(file_path: str) -> ExtractedDocument:
    """Ekstrak data dari file CSV."""
    def read_csv():
        import csv
        rows = []
        try:
            with open(file_path, newline="", encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                for i, row in enumerate(reader):
                    if i >= XLSX_MAX_ROWS:
                        rows.append(f"... (terpotong setelah {XLSX_MAX_ROWS} baris)")
                        break
                    rows.append("\t".join(row))
        except Exception as e:
            return f"[Gagal baca CSV: {e}]"
        return "\n".join(rows)

    text = await asyncio.to_thread(read_csv)
    page = ExtractedPage(page_num=0, text=text, is_scan=False)

    return ExtractedDocument(
        file_path=file_path,
        file_type="csv",
        total_pages=1,
        is_scanned=False,
        full_text=text,
        pages=[page],
    )


# ─── Image Extractor ─────────────────────────────────────────────────────────
async def _extract_image(file_path: str) -> ExtractedDocument:
    """OCR gambar menggunakan pytesseract + encode base64."""
    def read_image():
        import io
        try:
            import pytesseract
            from PIL import Image
            img = Image.open(file_path)
            text = pytesseract.image_to_string(img, lang="ind+eng").strip()
        except ImportError:
            text = "[pytesseract tidak terinstal]"
        except Exception as e:
            text = f"[Gagal OCR gambar: {e}]"

        # Encode base64
        try:
            with open(file_path, "rb") as f:
                b64 = base64.b64encode(f.read()).decode("utf-8")
        except Exception:
            b64 = ""
        return text, b64

    text, b64 = await asyncio.to_thread(read_image)
    page = ExtractedPage(page_num=0, text=text, is_scan=True, base64_image=b64)

    return ExtractedDocument(
        file_path=file_path,
        file_type="image",
        total_pages=1,
        is_scanned=True,
        full_text=text,
        pages=[page],
        base64_images=[b64] if b64 else [],
    )


# ─── Plain Text / Code Extractor ─────────────────────────────────────────────
async def _extract_plain_text(file_path: str) -> ExtractedDocument:
    """Baca file teks biasa (TXT, py, js, md, json, dll)."""
    def read_txt():
        try:
            return Path(file_path).read_text(encoding="utf-8", errors="replace")
        except Exception as e:
            return f"[Gagal baca file: {e}]"

    text = await asyncio.to_thread(read_txt)
    page = ExtractedPage(page_num=0, text=text, is_scan=False)

    return ExtractedDocument(
        file_path=file_path,
        file_type="text",
        total_pages=1,
        is_scanned=False,
        full_text=text,
        pages=[page],
    )


# ─── Cache Management ─────────────────────────────────────────────────────────
def clear_cache(cache_key: Optional[str] = None) -> None:
    """Hapus cache spesifik atau seluruh L1 RAM cache."""
    global _EXTRACTOR_CACHE
    if cache_key:
        _EXTRACTOR_CACHE.pop(cache_key, None)
    else:
        _EXTRACTOR_CACHE.clear()
    logger.info(f"[EXTRACTOR] Cache dihapus: {cache_key or 'semua'}")


def get_cache_stats() -> Dict[str, Any]:
    """Info diagnostik cache."""
    return {
        "l1_entries": len(_EXTRACTOR_CACHE),
        "l1_keys": list(_EXTRACTOR_CACHE.keys()),
        "l2_cache_dir": str(CACHE_DIR),
        "l2_files": len(list(CACHE_DIR.glob("*.extract.json"))),
    }
